"""Tests du logiciel de décorticage.

Lancer :  python -m unittest discover -s tests -v
Test d'intégration sur de vrais PDF (facultatif) :
    DECORTICAGE_PDF_DIR=/chemin/vers/pdfs python -m unittest discover -s tests
"""

from __future__ import annotations

import datetime as dt
import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from decorticage import securite
from decorticage.config import Identite, Reglages
from decorticage.pedagogie import (construire_seances, niveau_bloom, objectif_depuis_savoir, repartir,
                                   simplifier)
from decorticage.planification import Parametres, construire_plan, jours_feries_reunion, paques, _fenetres
from decorticage.programme import criteres_en_clair
from decorticage.publication import PublicationRefusee, RegistrePublication
from decorticage.referentiels import analyser_reac, detecter_type, joindre_items

REAC_MINI = """REFERENTIEL EMPLOI ACTIVITES COMPETENCES
DU TITRE PROFESSIONNEL
Vendeur test
Niveau 4
FICHE ACTIVITÉ TYPE N° 1
Vendre en magasin
Définition, description de l’activité type et conditions d’exercice
Le vendeur accueille et conseille.
Liste des compétences professionnelles de l’activité type
Accueillir le client
Conseiller le client et conclure la vente
Compétences transversales de l’activité type
Communiquer
FICHE COMPÉTENCE PROFESSIONNELLE N° 1
Accueillir le client
Description de la compétence – processus de mise en œuvre
Accueillir le client avec courtoisie. Identifier les besoins du client par le questionnement.
Contexte(s) professionnel(s) de mise en œuvre
En magasin.
Critères de performance
L’accueil est réalisé selon les standards de l’enseigne
Les besoins du client sont identifiés
Savoir-faire techniques, savoir-faire organisationnels, savoir-faire relationnels, savoirs
1 - Savoir-faire techniques
Utiliser les techniques de questionnement pour identifier les besoins du
client
2 - Savoir-faire organisationnels
Organiser l’espace d’accueil
3 - Savoir-faire relationnels
Pratiquer l’écoute active
4 - Savoirs
Connaissance des techniques de questionnement
FICHE COMPÉTENCE PROFESSIONNELLE N° 2
Conseiller le client et conclure la vente
Description de la compétence – processus de mise en œuvre
Présenter les produits. Conclure la vente.
Contexte(s) professionnel(s) de mise en œuvre
En magasin.
Critères de performance
L’argumentaire est adapté
Savoir-faire techniques, savoir-faire organisationnels, savoir-faire relationnels, savoirs
1 - Savoir-faire techniques
Construire un argumentaire de vente
2 - Savoir-faire organisationnels
Préparer les produits à présenter
3 - Savoir-faire relationnels
Traiter les objections du client
4 – Savoirs
Connaissance des techniques de vente
FICHE DES COMPÉTENCES TRANSVERSALES DE L’EMPLOI TYPE
Communiquer
"""


class TestSecurite(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def test_nom_fichier(self):
        self.assertEqual(securite.nom_fichier_sur("../../etc/passwd"), "passwd")
        self.assertEqual(securite.nom_fichier_sur("..\\..\\x.pdf"), "x.pdf")
        self.assertNotIn("/", securite.nom_fichier_sur("a/b/c d.pdf"))

    def test_chemin_dans(self):
        with self.assertRaises(securite.FichierRefuse):
            securite.chemin_dans(self.dir, "..", "secret")
        self.assertTrue(str(securite.chemin_dans(self.dir, "ok.txt")).startswith(str(self.dir.resolve())))

    def test_pdf_deguise(self):
        f = self.dir / "faux.pdf"
        f.write_bytes(b"MZ\x90\x00" * 20)
        with self.assertRaises(securite.FichierRefuse):
            securite.verifier_fichier(f)

    def test_formats_refuses(self):
        for ext in (".xlsm", ".docm", ".doc", ".exe"):
            f = self.dir / f"x{ext}"
            f.write_bytes(b"PK\x03\x04data")
            with self.assertRaises(securite.FichierRefuse):
                securite.verifier_fichier(f)

    def _zip(self, nom, entrees):
        f = self.dir / nom
        with zipfile.ZipFile(f, "w", zipfile.ZIP_DEFLATED) as z:
            for n, d in entrees.items():
                z.writestr(n, d)
        return f

    def test_macro_cachee_dans_docx(self):
        f = self._zip("a.docx", {"word/document.xml": "<x/>", "word/vbaProject.bin": "x"})
        with self.assertRaises(securite.FichierRefuse):
            securite.verifier_fichier(f)

    def test_bombe_zip(self):
        f = self._zip("b.xlsx", {"xl/workbook.xml": "<x/>", "xl/gros.xml": "0" * 5_000_000})
        with self.assertRaises(securite.FichierRefuse):
            securite.verifier_fichier(f)

    def test_chemin_piege_dans_zip(self):
        f = self._zip("c.docx", {"word/document.xml": "<x/>", "../../evil.txt": "x"})
        with self.assertRaises(securite.FichierRefuse):
            securite.verifier_fichier(f)

    def test_formule_neutralisee(self):
        from openpyxl import Workbook
        ws = Workbook().active
        securite.ecrire_cellule(ws["A1"], '=HYPERLINK("http://evil","clic")')
        self.assertEqual(ws["A1"].data_type, "s")
        self.assertEqual(securite.texte_csv("=1+1"), "'=1+1")

    def test_mot_de_passe(self):
        ancien = securite.ITERATIONS
        securite.ITERATIONS = 1000
        try:
            c = securite.CoffreAcces(self.dir)
            with self.assertRaises(ValueError):
                c.definir("court")
            c.definir("Tres-Long-Mot2Passe")
            self.assertTrue(c.verifier("Tres-Long-Mot2Passe"))
            for _ in range(securite.ECHECS_AVANT_BLOCAGE):
                self.assertFalse(c.verifier("mauvais"))
            self.assertGreater(c.attente_restante(), 0)
            self.assertFalse(c.verifier("Tres-Long-Mot2Passe"))  # bloqué même avec le bon
        finally:
            securite.ITERATIONS = ancien


class TestReferentiels(unittest.TestCase):
    def test_joindre_items(self):
        items = joindre_items(["Utiliser les outils de", "planification (GANTT,", "PERT)", "Animer une réunion"])
        self.assertEqual(items, ["Utiliser les outils de planification (GANTT, PERT)", "Animer une réunion"])

    def test_reac(self):
        self.assertEqual(detecter_type(REAC_MINI), "REAC")
        r = analyser_reac(REAC_MINI)
        self.assertEqual(len(r.activites), 1)
        self.assertEqual(r.activites[0].competences, [1, 2])
        cp1 = r.competence(1)
        self.assertEqual(cp1.intitule, "Accueillir le client")
        self.assertEqual(len(cp1.criteres), 2)
        self.assertEqual(cp1.sf_techniques, ["Utiliser les techniques de questionnement pour identifier les besoins du client"])
        self.assertEqual(r.competence(2).savoirs, ["Connaissance des techniques de vente"])
        self.assertEqual(r.avertissements, [])


class TestPedagogie(unittest.TestCase):
    def test_bloom(self):
        self.assertEqual(niveau_bloom("Analyser les écarts"), 4)
        self.assertEqual(niveau_bloom("Élaborer un plan d’actions"), 6)
        self.assertEqual(niveau_bloom("Contrôler la conformité"), 5)

    def test_savoir(self):
        self.assertEqual(objectif_depuis_savoir("Connaissance des techniques d’animation de réunion"),
                         "Décrire les techniques d’animation de réunion")

    def test_simplifier(self):
        self.assertEqual(simplifier("Animer l’équipe de son périmètre"), "Animer l’équipe")

    def test_repartir_exact(self):
        for total in (600, 615, 4500):
            d = repartir(total, [3, 1, 1.5, 2])
            self.assertEqual(sum(d), total - total % 15)
            self.assertTrue(all(x % 15 == 0 for x in d))

    def test_aucun_savoir_perdu(self):
        r = analyser_reac(REAC_MINI)
        for cp in r.competences:
            seances = construire_seances(cp)
            contenus = [c for s in seances for o in s.objectifs for c in o.contenus]
            self.assertEqual(len(contenus), len(cp.savoirs))


class TestPlanification(unittest.TestCase):
    def test_paques_et_feries(self):
        self.assertEqual(paques(2027), dt.date(2027, 3, 28))
        f = jours_feries_reunion(2026)
        self.assertIn(dt.date(2026, 12, 20), f)
        self.assertIn(dt.date(2026, 4, 6), f)  # lundi de Pâques 2026

    def test_pauses_exclues(self):
        p = Parametres(date_debut=dt.date(2027, 1, 4), heures_centre=10)
        self.assertEqual(sum(b - a for a, b in _fenetres(p)), 390)

    def test_plan_complet(self):
        r = analyser_reac(REAC_MINI)
        p = Parametres(date_debut=dt.date(2026, 12, 14), heures_centre=70, h_accueil=3.5, h_ecf_par_at=3.5,
                       h_dossier_pro=3.5, h_productions=0, h_examen_blanc=3.5, h_bilan=3.5)
        plan = construire_plan(r, None, p)
        self.assertEqual(sum(l["minutes"] for l in plan.lignes), 70 * 60)
        jours = {l["date"] for l in plan.lignes}
        self.assertNotIn(dt.date(2026, 12, 20), jours)
        self.assertNotIn(dt.date(2026, 12, 25), jours)
        self.assertTrue(all(l["date"].weekday() < 5 for l in plan.lignes))

    def test_rythme_alternance(self):
        r = analyser_reac(REAC_MINI)
        p = Parametres(date_debut=dt.date(2027, 1, 4), heures_centre=70, rythme="CE")
        plan = construire_plan(r, None, p)
        semaines = {l["date"].isocalendar()[1] for l in plan.lignes}
        self.assertTrue(all(s % 2 == 1 for s in semaines))  # 1 semaine sur 2 depuis la S1


class TestProgrammeEtPublication(unittest.TestCase):
    def test_criteres_en_clair(self):
        self.assertEqual(criteres_en_clair("Les besoins humains sont anticipés et adaptés à l’activité prévue"),
                         "Anticiper et adapter les besoins humains à l’activité prévue")
        self.assertEqual(criteres_en_clair("Les anomalies de livraison sont repérées, signalées et traitées suivant le contexte"),
                         "Repérer, signaler et traiter les anomalies de livraison suivant le contexte")

    def test_commanditaire_masque_yeba(self):
        self.assertTrue(Identite().est_yeba)
        i = Identite(commanditaire="ICCI")
        self.assertFalse(i.est_yeba)
        self.assertFalse(i.logo_autorise())

    def test_publication_trois_choix(self):
        reg = RegistrePublication(Path(tempfile.mkdtemp()))
        self.assertTrue(reg.doit_demander("RNCP1"))
        reg.decider("RNCP1", "non_redemander")
        self.assertTrue(reg.doit_demander("RNCP1"))
        with self.assertRaises(PublicationRefusee):
            reg.verifier_publiable("RNCP1")
        reg.decider("RNCP1", "oui")
        reg.verifier_publiable("RNCP1")
        reg.decider("RNCP1", "non_definitif")
        self.assertFalse(reg.doit_demander("RNCP1"))
        with self.assertRaises(PublicationRefusee):
            reg.decider("RNCP1", "oui")

    def test_reglages_formulaire(self):
        r = Reglages.depuis_dict({"heures_centre": "455,5", "afficher_logo": "1"})
        self.assertEqual(r.heures_centre, 455.5)
        self.assertTrue(r.afficher_logo)
        self.assertFalse(r.afficher_nom)  # case décochée


@unittest.skipUnless(os.environ.get("DECORTICAGE_PDF_DIR"), "PDF réels non fournis")
class TestIntegrationPDF(unittest.TestCase):
    def test_bout_en_bout(self):
        from openpyxl import load_workbook

        from decorticage.moteur import analyser_documents, construire_contexte, generer_sorties
        d = Path(os.environ["DECORTICAGE_PDF_DIR"])
        os.environ["YEBA_DATA_DIR"] = tempfile.mkdtemp()
        analyse = analyser_documents(sorted(d.glob("*.pdf")))
        ctx = construire_contexte(analyse, Reglages(heures_centre=455, date_debut="2027-01-04"))
        sorties = generer_sorties(ctx, Path(tempfile.mkdtemp()))
        wb = load_workbook(sorties["decorticage"])
        self.assertIn("Décorticage", wb.sheetnames)
        self.assertNotIn("openpyxl", (wb.properties.creator or "").lower())
        for ws in wb.worksheets:
            self.assertEqual(ws.page_setup.orientation, "landscape")
        couverture = wb["Couverture critères"]
        self.assertFalse(any("À COUVRIR" in str(c.value) for row in couverture.iter_rows() for c in row))


if __name__ == "__main__":
    unittest.main()
