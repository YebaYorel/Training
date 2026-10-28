"""Exports Word (programme, examen), page web (site) et agenda (.ics)."""

from __future__ import annotations

import datetime as dt
import html
import json
import re
from collections import OrderedDict
from pathlib import Path

from .planification import fmt_hm
from .programme import construire_examen, construire_programme


# ==========================================================================
# Word (modifiable, paysage, gros caractères)
# ==========================================================================
def _doc_base(ctx):
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt, RGBColor

    d = docx.Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.6))
    st = d.styles["Normal"]
    st.font.name = ctx.charte.police
    st.font.size = Pt(ctx.charte.taille_texte)
    for niveau, taille in ((1, ctx.charte.taille_titre), (2, ctx.charte.taille_texte + 4)):
        h = d.styles[f"Heading {niveau}"]
        h.font.name = ctx.charte.police
        h.font.size = Pt(taille)
        h.font.color.rgb = RGBColor.from_string(ctx.charte.couleur_primaire if niveau == 1 else ctx.charte.couleur_secondaire)
    props = d.core_properties
    nom = ctx.identite.nom_affiche(ctx.charte)
    props.author = nom
    props.last_modified_by = nom
    props.comments = ""
    logo = ctx.charte.chemin_logo(ctx.base) if ctx.identite.logo_autorise() else None
    if logo:
        d.add_picture(str(logo), height=Cm(2.2))   # image indépendante : supprimable d'un clic
    return d


def _tableau_docx(d, ctx, lignes: list[list[str]]):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    from docx.shared import Cm

    t = d.add_table(rows=0, cols=len(lignes[0]))
    t.style = "Table Grid"
    t.autofit = False
    # largeurs proportionnelles au contenu (bornées) sur 26,5 cm utiles en paysage
    poids = []
    for j in range(len(lignes[0])):
        longueurs = [len(str(r[j])) for r in lignes if j < len(r)]
        mini = max(12, len(str(lignes[0][j])) + 4)   # l'en-tête ne doit jamais couper un mot
        poids.append(min(60, max(mini, sum(longueurs) / max(len(longueurs), 1))))
    largeurs = [Cm(26.5 * w / sum(poids)) for w in poids]
    for i, row in enumerate(lignes):
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].width = largeurs[j]
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            run = p.add_run(str(v))
            run.font.size = Pt(ctx.charte.taille_texte)
            if i == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), ctx.charte.couleur_primaire)
                cells[j]._tc.get_or_add_tcPr().append(shd)
    for j, col in enumerate(t.columns):   # grille du tableau (lue par Word et LibreOffice)
        col.width = largeurs[j]
    # répéter la ligne d'en-tête à chaque page
    trPr = t.rows[0]._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)
    d.add_paragraph()


def exporter_programme_docx(ctx, chemin: Path) -> Path:
    intitule, sections = construire_programme(ctx.reac, ctx.re, ctx.fiche, ctx.plan, ctx.reglages, ctx.tarif)
    d = _doc_base(ctx)
    d.add_heading(f"Programme de formation — {intitule}", 1)
    nom = ctx.identite.nom_affiche(ctx.charte)
    if nom:
        d.add_paragraph(nom)
    for s in sections:
        d.add_heading(s.titre, 2)
        for p in s.paragraphes:
            d.add_paragraph(p)
        for p in s.puces:
            d.add_paragraph(p, style="List Bullet")
        if s.tableau:
            _tableau_docx(d, ctx, s.tableau)
    d.add_paragraph(f"Mise à jour : {dt.date.today():%d/%m/%Y}")
    d.save(str(chemin))
    return chemin


def exporter_examen_docx(ctx, chemin: Path) -> Path:
    ex = construire_examen(ctx.reac, ctx.re, ctx.fiche)
    d = _doc_base(ctx)
    titre = ctx.fiche.intitule if ctx.fiche and ctx.fiche.intitule else ctx.reac.intitule
    d.add_heading(f"L'examen en clair — {titre}", 1)
    if ex["duree_totale"]:
        d.add_paragraph(f"Durée totale pour le candidat : {ex['duree_totale']}")
    d.add_heading("Partie 1 — Les épreuves", 2)
    if len(ex["partie1"]) > 1:
        _tableau_docx(d, ctx, ex["partie1"])
    for c in ex["consignes"]:
        d.add_paragraph(c, style="List Bullet")
    d.add_heading("Partie 2 — Ce que le jury attend", 2)
    for b in ex["partie2"]:
        d.add_heading(b["cp"], 2)
        if b["epreuves"]:
            d.add_paragraph("Évaluée par : " + ", ".join(b["epreuves"]))
        for a in b["attendus"]:
            d.add_paragraph(a, style="List Bullet")
        if b["production"]:
            d.add_paragraph(f"À préparer avant l'examen — partie de document ({b['production']['pages']}) :")
            for t in b["production"]["themes"]:
                d.add_paragraph(t, style="List Bullet 2")
    d.save(str(chemin))
    return chemin


# ==========================================================================
# Page web du programme (autonome, sans ressource externe)
# ==========================================================================
def slug(texte: str) -> str:
    import unicodedata
    s = unicodedata.normalize("NFKD", texte).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")[:80] or "programme"


def page_programme_html(ctx) -> tuple[str, dict]:
    intitule, sections = construire_programme(ctx.reac, ctx.re, ctx.fiche, ctx.plan, ctx.reglages, ctx.tarif)
    e = html.escape
    c = ctx.charte
    nom = ctx.identite.nom_affiche(c)
    corps = []
    for s in sections:
        corps.append(f"<section><h2>{e(s.titre)}</h2>")
        corps += [f"<p>{e(p)}</p>" for p in s.paragraphes]
        if s.puces:
            corps.append("<ul>" + "".join(f"<li>{e(p)}</li>" for p in s.puces) + "</ul>")
        if s.tableau:
            tete = "".join(f"<th scope='col'>{e(x)}</th>" for x in s.tableau[0])
            lignes = "".join("<tr>" + "".join(f"<td>{e(str(x))}</td>" for x in r) + "</tr>" for r in s.tableau[1:])
            corps.append(f"<div class='t'><table><thead><tr>{tete}</tr></thead><tbody>{lignes}</tbody></table></div>")
        corps.append("</section>")
    donnees = {
        "@context": "https://schema.org", "@type": "Course", "name": intitule,
        "courseCode": ctx.fiche.code if ctx.fiche else "",
        "educationalCredentialAwarded": intitule,
        "provider": {"@type": "Organization", "name": nom} if nom else None,
        "inLanguage": "fr",
    }
    donnees = {k: v for k, v in donnees.items() if v}
    jsonld = json.dumps(donnees, ensure_ascii=False).replace("</", "<\\/")
    page = f"""<!doctype html>
<html lang="fr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; img-src 'self' data:">
<title>{e(intitule)}</title>
<script type="application/ld+json">{jsonld}</script>
<style>
:root{{--p:#{c.couleur_primaire};--s:#{c.couleur_secondaire};--c:#{c.couleur_claire};--t:#{c.couleur_texte}}}
body{{font-family:{e(c.police)},Arial,sans-serif;font-size:1.25rem;line-height:1.6;color:var(--t);margin:0;background:#fff}}
header{{background:var(--p);color:#fff;padding:1.5rem 1rem}}
header h1{{margin:0;font-size:2rem}}
main{{max-width:1100px;margin:0 auto;padding:1rem}}
h2{{color:var(--p);border-left:.4rem solid var(--s);padding-left:.6rem;margin-top:2rem}}
.t{{overflow-x:auto}}
table{{border-collapse:collapse;width:100%;font-size:1.1rem}}
th{{background:var(--p);color:#fff;text-align:left}}
th,td{{border:1px solid #9aa5b1;padding:.5rem;vertical-align:top}}
tr:nth-child(even) td{{background:var(--c)}}
a{{color:var(--p)}}
</style>
</head>
<body>
<header><h1>{e(intitule)}</h1>{f"<p>{e(nom)}</p>" if nom else ""}</header>
<main>
{''.join(corps)}
<p><small>Mise à jour : {dt.date.today():%d/%m/%Y}</small></p>
</main>
</body>
</html>
"""
    return page, {"intitule": intitule, "code": ctx.fiche.code if ctx.fiche else "",
                  "sections": [{"titre": s.titre, "paragraphes": s.paragraphes, "puces": s.puces,
                                "tableau": s.tableau} for s in sections]}


# ==========================================================================
# Agenda (.ics) — importable dans n'importe quel calendrier, fichier local
# ==========================================================================
def _ics_texte(s: str) -> str:
    s = s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")
    return re.sub(r"[\x00-\x1f]", "", s)


def _pli(ligne: str) -> str:
    """Pliage RFC 5545 à 75 octets."""
    sortie, courant = [], ""
    for ch in ligne:
        if len((courant + ch).encode("utf-8")) > 74:
            sortie.append(courant)
            courant = " " + ch
        else:
            courant += ch
    sortie.append(courant)
    return "\r\n".join(sortie)


def exporter_ics(ctx, chemin: Path) -> Path:
    lignes = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Decorticage//FR", "CALSCALE:GREGORIAN",
              "BEGIN:VTIMEZONE", "TZID:Indian/Reunion", "BEGIN:STANDARD", "DTSTART:19700101T000000",
              "TZOFFSETFROM:+0400", "TZOFFSETTO:+0400", "TZNAME:RET", "END:STANDARD", "END:VTIMEZONE"]
    groupes: "OrderedDict[tuple, list]" = OrderedDict()
    for l in ctx.plan.lignes:
        groupes.setdefault((l["date"], l["seance"].numero), []).append(l)
    horodatage = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    for (jour, num), ls in groupes.items():
        s = ls[0]["seance"]
        debut, fin = ls[0]["debut"], ls[-1]["fin"]
        objs = "\n".join(dict.fromkeys(f"- {x['objectif'].texte}" for x in ls))
        lignes += [
            "BEGIN:VEVENT",
            f"UID:{slug(ctx.titre)}-{jour:%Y%m%d}-s{num}-{debut}@decorticage.local",
            f"DTSTAMP:{horodatage}",
            f"DTSTART;TZID=Indian/Reunion:{jour:%Y%m%d}T{debut // 60:02d}{debut % 60:02d}00",
            f"DTEND;TZID=Indian/Reunion:{jour:%Y%m%d}T{fin // 60:02d}{fin % 60:02d}00",
            _pli(f"SUMMARY:{_ics_texte(f'S{num} — {s.titre}')}"),
            _pli(f"DESCRIPTION:{_ics_texte(f'{fmt_hm(debut)}-{fmt_hm(fin)}' + chr(10) + objs)}"),
            "END:VEVENT",
        ]
    lignes.append("END:VCALENDAR")
    chemin.write_text("\r\n".join(lignes) + "\r\n", encoding="utf-8")
    return chemin
