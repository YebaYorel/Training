"""Lecture du texte des documents reçus (PDF, Word, Excel, texte).

Tout est lu localement, sans service en ligne. Chaque format passe d'abord
par `securite.verifier_fichier`.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from .securite import NB_MAX_PAGES_PDF, FichierRefuse, verifier_fichier

# En-têtes/pieds de page des référentiels du ministère, à retirer du texte.
_PIED_PAGE = re.compile(
    r"^\s*(SIGLE\s+Type de document.*|[A-Z]{2,6}\s+(REAC|RE|RC)\s+TP-\d{5}\s+\d+.*|"
    r"https?://www\.francecompetences\.fr/\S+\s+\d{2}/\d{2}/\d{4}.*|Page \d+ sur \d+)\s*$"
)


def _nettoyer(texte: str) -> str:
    lignes = []
    for ligne in texte.splitlines():
        if _PIED_PAGE.match(ligne):
            continue
        # puces des PDF (Symbol/Wingdings) → tiret standard
        ligne = re.sub(r"^\s*[•▪●]+\s*", "- ", ligne)
        lignes.append(ligne.rstrip())
    texte = "\n".join(lignes)
    texte = texte.replace("’", "’").replace(" ", " ")
    return re.sub(r"\n{3,}", "\n\n", texte)


def _lire_pdf(chemin: Path) -> str:
    import pdfplumber  # import local : chargé seulement si besoin

    morceaux = []
    with pdfplumber.open(str(chemin)) as pdf:
        if len(pdf.pages) > NB_MAX_PAGES_PDF:
            raise FichierRefuse(f"PDF trop long (plus de {NB_MAX_PAGES_PDF} pages).")
        for page in pdf.pages:
            morceaux.append(page.extract_text() or "")
    texte = "\n".join(morceaux)
    if len(texte.strip()) < 50:
        raise FichierRefuse(
            "Ce PDF ne contient pas de texte lisible (probablement un scan). "
            "Demandez la version numérique officielle (site travail-emploi.gouv.fr "
            "ou France Compétences)."
        )
    return texte


def _lire_docx(chemin: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(chemin))
    lignes = [p.text for p in document.paragraphs]
    for table in document.tables:
        for rangee in table.rows:
            cellules = []
            for c in rangee.cells:
                t = c.text.strip()
                if t and (not cellules or cellules[-1] != t):
                    cellules.append(t)
            lignes.append(" | ".join(cellules))
    return "\n".join(lignes)


def _lire_xlsx(chemin: Path) -> str:
    from openpyxl import load_workbook

    # read_only + data_only : on ne lit que les valeurs, jamais les formules.
    wb = load_workbook(str(chemin), read_only=True, data_only=True)
    lignes = []
    try:
        for ws in wb.worksheets:
            for rangee in ws.iter_rows(values_only=True):
                valeurs = [str(v).strip() for v in rangee if v is not None and str(v).strip()]
                if valeurs:
                    lignes.append(" | ".join(valeurs))
    finally:
        wb.close()
    return "\n".join(lignes)


def _lire_texte(chemin: Path) -> str:
    brut = chemin.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return brut.decode(enc)
        except UnicodeDecodeError:
            continue
    return brut.decode("utf-8", errors="replace")


def lire_document(chemin: Path, nettoyer: bool = True) -> str:
    """Retourne le texte d'un document, après contrôle de sécurité.

    `nettoyer=False` conserve les cartouches (sigle, millésime, dates).
    """
    chemin = Path(chemin)
    ext = verifier_fichier(chemin)
    if ext == ".pdf":
        texte = _lire_pdf(chemin)
    elif ext == ".docx":
        texte = _lire_docx(chemin)
    elif ext == ".xlsx":
        texte = _lire_xlsx(chemin)
    else:
        texte = _lire_texte(chemin)
    return _nettoyer(texte) if nettoyer else texte


def lire_tableau(chemin: Path) -> list[dict]:
    """Lit un tableau (CSV ou XLSX) en liste de dictionnaires (1re ligne = titres)."""
    chemin = Path(chemin)
    ext = verifier_fichier(chemin, {".csv", ".xlsx"})
    if ext == ".csv":
        texte = _lire_texte(chemin)
        echantillon = texte[:4096]
        try:
            dialecte = csv.Sniffer().sniff(echantillon, delimiters=";,\t")
        except csv.Error:
            dialecte = csv.excel
            dialecte.delimiter = ";"
        return [dict(r) for r in csv.DictReader(io.StringIO(texte), dialect=dialecte)]

    from openpyxl import load_workbook

    wb = load_workbook(str(chemin), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rangees = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rangees:
        return []
    titres = [str(t).strip() if t is not None else f"colonne_{i}" for i, t in enumerate(rangees[0])]
    return [
        {titres[i]: v for i, v in enumerate(r) if i < len(titres)}
        for r in rangees[1:]
        if any(v is not None for v in r)
    ]
