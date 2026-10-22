"""Pilotage de l'année : plan d'actions daté, pièces par financeur, budget et trésorerie, tableau de bord."""
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from calendrier import PROMOTIONS, fmt, fmt_court, jour_off, periodes, planifier, session_examen
from lib import HORAIRES, ORGANISME, SORTIE, Doc
from titres import DUREE_JOUR, heures_centre

O, H = ORGANISME, HORAIRES
DOSSIER = "05_Pilotage_annuel"


def _jour_heure(code, heure_cible):
    """Date de la journée où le cumul d'heures atteint heure_cible."""
    jours, _ = planifier(code)
    cumul = 0
    for d, lieu, *_ in jours:
        if lieu == "Centre":
            cumul += DUREE_JOUR
            if cumul >= heure_cible:
                return d


def _moins_mois(d, n):
    m = d.month - n
    a = d.year + (m - 1) // 12
    m = (m - 1) % 12 + 1
    return dt.date(a, m, min(d.day, 28))


def jalons():
    """Plan d'actions daté de l'organisme, de novembre 2026 à la fin de la promotion RPMS."""
    out = [
        (dt.date(2026, 10, 30), "DEETS", "Envoyer la réponse version A à la DEETS (date limite pratique : vendredi 30 octobre 2026)"),
        (dt.date(2026, 11, 6), "Qualité", "Vérifier le périmètre et la date d'audit de surveillance du certificat Qualiopi 25FOR02027.1"),
        (dt.date(2026, 11, 13), "Référentiels", "Obtenir le REAC et le référentiel d'évaluation RPMS millésime 08 (DEETS ou ministère du Travail)"),
        (dt.date(2026, 11, 20), "Examen", "Décider : agrément de centre d'examen ou partenariat avec un centre agréé (document AGR-03)"),
        (dt.date(2026, 11, 27), "Locaux", "Réserver au Secrétariat Center les journées de la promotion MEM ; vérifier accessibilité, sanitaires, salle de pause"),
        (dt.date(2026, 11, 30), "Sécurité", "Plan de continuité cyclonique : contacts, formation à distance, procédure d'information des stagiaires"),
    ]
    for code, p in PROMOTIONS.items():
        jours, _ = planifier(code)
        debut, fin = jours[0][0], jours[-1][0]
        s1, _ = session_examen(code)
        n = p["nom"]
        out += [
            (_moins_mois(s1, 7), "Examen", f"{n} : déposer la demande d'agrément de session (démarche en ligne DEETS 974) ou signer la convention avec le centre agréé"),
            (_moins_mois(debut, 3), "Offre", f"{n} : publier l'offre (Carif-Oref Réunion, KAIROS, Mon Compte Formation, site) ; informer Missions Locales et agences France Travail"),
            (_moins_mois(debut, 2), "Partenaires", f"{n} : constituer le vivier d'entreprises d'accueil pour les périodes en entreprise"),
            (_moins_mois(debut, 2), "Recrutement", f"{n} : 1re information collective"),
            (_moins_mois(debut, 1), "Recrutement", f"{n} : 2e information collective, entretiens, décisions d'admission, accords de financement"),
            (debut - dt.timedelta(days=15), "Administratif", f"{n} : envoyer convocations, livret d'accueil, règlement intérieur, programme"),
            (debut, "Formation", f"{n} : démarrage — attestations d'entrée, déclarations d'entrée sur les plateformes des financeurs"),
            (_jour_heure(code, 28), "Formation", f"{n} : élection des délégués des stagiaires (entre la 20e et la 40e heure)"),
        ]
        for i, (lieu, a, b, nb, _) in enumerate([x for x in periodes(code) if x[0] == "Entreprise"], 1):
            out += [(a - dt.timedelta(days=14), "Entreprise", f"{n} : conventions de la période en entreprise n° {i} signées"),
                    (a + dt.timedelta(days=(b - a).days // 2), "Entreprise", f"{n} : appel ou visite des tuteurs (période n° {i})")]
        out += [
            (s1 - dt.timedelta(days=15), "Examen", f"{n} : dossiers professionnels, productions et attestations de période en entreprise complets"),
            (fin, "Administratif", f"{n} : évaluation à chaud, attestations de fin de formation, certificats de réalisation, facturation"),
            (s1, "Examen", f"{n} : session de validation [À VÉRIFIER : date confirmée]"),
            (fin + dt.timedelta(days=91), "Qualité", f"{n} : questionnaire à froid (3 mois)"),
            (fin + dt.timedelta(days=182), "Qualité", f"{n} : enquête d'insertion à 6 mois et mise à jour des résultats publiés"),
        ]
    for annee in (2027, 2028):
        out += [(dt.date(annee, 5, 31), "Obligation", f"Déposer le bilan pédagogique et financier de {annee - 1} (Mon Activité Formation)"),
                (dt.date(annee, 3, 31), "Obligation", "Renouveler l'attestation de vigilance URSSAF (tous les 6 mois)"),
                (dt.date(annee, 9, 30), "Obligation", "Renouveler l'attestation de vigilance URSSAF (tous les 6 mois)"),
                (dt.date(annee, 4, 30), "RGPD", "Revue annuelle du registre des traitements et purge des données échues")]
    out.append((dt.date(2027, 1, 31), "IA Act", "Former l'équipe à la maîtrise de l'IA (article 4) et faire signer la charte IA"))
    return sorted((_ouvre(x), dom, txt) for x, dom, txt in out)


def _ouvre(d):
    """Reporte une échéance tombant un week-end ou un jour férié au jour ouvré suivant."""
    while jour_off(d):
        d += dt.timedelta(days=1)
    return d


def doc_plan():
    d = Doc("Plan d'actions de l'année", "Ce qu'il faut faire, et quand — novembre 2026 à 2028", code="PIL-01", paysage=True)
    d.p("Dates calculées à partir des calendriers des promotions. Cochez au fur et à mesure ; une ligne en retard se traite en priorité.")
    d.tableau(["Fait", "Date", "Domaine", "Action"],
              [["☐", fmt_court(x), dom, txt] for x, dom, txt in jalons()], [1.5, 3, 3.4, 17.8])
    return d.save(f"{DOSSIER}/PIL-01_Plan_actions_annee.docx")


SOCLE = [
    "Avis de situation INSEE ou extrait RNE de moins de 3 mois (entreprise individuelle)",
    f"Récépissé de déclaration d'activité (NDA {O['nda']})",
    f"Certificat Qualiopi en cours de validité ({O['qualiopi']}) et périmètre (catégories)",
    "Attestation de vigilance URSSAF de moins de 6 mois (obligatoire pour tout contrat d'au moins 5 000 € HT)",
    "Attestation de régularité fiscale (marchés publics)",
    "Attestation d'assurance responsabilité civile professionnelle",
    "RIB au nom de l'entreprise",
    "Dernier bilan pédagogique et financier déposé",
    "Attestation d'exonération de TVA [À VÉRIFIER]",
    "CV et justificatifs des formateurs ; contrats de sous-traitance",
    "Conditions de location du Secrétariat Center (lieu de formation)",
    "Programme, livret d'accueil, règlement intérieur, CGV, procédures handicap et réclamations",
]

FINANCEURS = [
    ("Caisse des Dépôts — CPF (EDOF)", [
        "Compte organisme sur EDOF ; preuve du droit à préparer la certification [À VÉRIFIER : pour un titre professionnel, agrément DEETS ou partenariat avec un centre agréé]",
        "Fiche action complète (programme, prix, prérequis, modalités, accessibilité)",
        "Déclarations d'entrée, de sortie et de service fait dans les délais",
        "Émargements et preuve de passage de la certification"]),
    ("France Travail (AIF, AFC, POEI, POEC)", [
        "Offre référencée au Carif-Oref Réunion", "Devis et attestations de présence sur KAIROS",
        "Programme, calendrier, conventions de période en entreprise", "Attestation de fin de formation, résultats de certification"]),
    ("Missions Locales", [
        "Fiche de présentation du centre et calendrier des informations collectives",
        "Fiche de liaison prescripteur (entrée, mi-parcours, sortie) avec l'accord du jeune"]),
    ("Région Réunion", [
        "Réponse aux appels d'offres du programme régional de formation : DC1, DC2 ou DUME, attestations, références, mémoire technique",
        "Bilans pédagogiques, émargements, suivi de rémunération des stagiaires, indicateurs d'insertion",
        "Cofinancement FSE+ éventuel : publicité européenne et conservation renforcée [À VÉRIFIER]"]),
    ("OPCO (entreprises)", [
        "Convention, programme, devis ; demande de prise en charge déposée par l'entreprise avant le démarrage",
        "Émargements, certificat de réalisation, facture"]),
    ("DEETS de La Réunion", [
        "Déclaration d'activité à jour (toute modification signalée)", "Bilan pédagogique et financier annuel",
        "Agrément des sessions de titre professionnel ou convention avec un centre agréé",
        "Règlement intérieur et élection des délégués (actions de plus de 500 h)"]),
    ("URSSAF", [
        "Déclarations sociales et contribution à la formation professionnelle",
        "Attestation de vigilance tous les 6 mois ; vérification de celle des formateurs indépendants"]),
    ("Auditeur Qualiopi", ["Preuves des indicateurs (voir la matrice QUA-02)",
                           "Audit de surveillance entre le 14e et le 22e mois ; renouvellement à 3 ans"]),
]


def doc_pieces():
    d = Doc("Pièces à fournir aux financeurs et contrôleurs", "OPCO • DEETS • France Travail • Missions Locales • URSSAF • Région • Caisse des Dépôts", code="PIL-02")
    d.h1("1. Socle commun, à tenir toujours à jour")
    d.cases(SOCLE)
    d.h1("2. Pièces par interlocuteur")
    for nom, items in FINANCEURS:
        d.h2(nom)
        d.cases(items)
    return d.save(f"{DOSSIER}/PIL-02_Pieces_financeurs_controleurs.docx")


# ---------------------------------------------------------------- Excel
BLEU_F = PatternFill("solid", fgColor="1B3A6B")
SAISIE_F = PatternFill("solid", fgColor="FBF6E9")
CALC_F = PatternFill("solid", fgColor="E8EEF7")
BLANC = Font(color="FFFFFF", bold=True, size=12, name="Arial")
NORMAL = Font(size=12, name="Arial")
GRAS = Font(size=12, name="Arial", bold=True, color="1B3A6B")
FIN = Side(style="thin", color="BFBFBF")
BORD = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)
EUROS = '#,##0 "€"'


def _entete(ws, ligne, valeurs):
    for i, v in enumerate(valeurs, 1):
        c = ws.cell(row=ligne, column=i, value=v)
        c.fill, c.font, c.border = BLEU_F, BLANC, BORD
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _cellule(ws, ref, valeur, fill=None, fmt=None, gras=False):
    c = ws[ref]
    c.value = valeur
    c.font = GRAS if gras else NORMAL
    c.border = BORD
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt


def _budget(wb, code):
    p = PROMOTIONS[code]
    t = p["titre"]
    jours, _ = planifier(code)
    nb_centre = sum(1 for j in jours if j[1] == "Centre")
    ws = wb.create_sheet(f"Budget {code}")
    ws["A1"] = f"Budget prévisionnel — {p['nom']} ({fmt_court(jours[0][0])} → {fmt_court(jours[-1][0])})"
    ws["A1"].font = Font(size=14, bold=True, name="Arial", color="1B3A6B")
    ws["A2"] = "Cases beiges = à saisir. Cases bleues = calculées."
    ws["A2"].font = Font(size=11, italic=True, name="Arial")
    _entete(ws, 4, ["Hypothèses", "Valeur", "Commentaire"])
    hyp = [
        ("Nombre de stagiaires", 10, "Saisir l'effectif prévu", SAISIE_F, "0"),
        ("Prix du parcours par stagiaire", None, "[À COMPLÉTER] — prix de vente, net de TVA", SAISIE_F, EUROS),
        ("Heures en centre", heures_centre(t), "Calculé à partir du déroulé", CALC_F, "0"),
        ("Journées en centre", nb_centre, "Calculé à partir du calendrier", CALC_F, "0"),
        ("Location de salle par jour", H["location_jour"], "Secrétariat Center", SAISIE_F, EUROS),
        ("Heures assurées par des formateurs externes", 0, "Heures non assurées par le dirigeant", SAISIE_F, "0"),
        ("Coût horaire d'un formateur externe", None, "[À COMPLÉTER]", SAISIE_F, EUROS),
        ("Ordinateurs portables (achat ou location)", None, "Indispensables pour T2 et l'épreuve de mise en situation du MEM", SAISIE_F, EUROS),
        ("Coût d'examen par candidat (centre agréé partenaire)", None, "[À COMPLÉTER] si pas d'agrément propre", SAISIE_F, EUROS),
        ("Jury blanc (défraiement)", None, "Professionnels invités", SAISIE_F, EUROS),
        ("Communication et recrutement", None, "Informations collectives, supports", SAISIE_F, EUROS),
        ("Assurances, logiciels, frais administratifs", None, "Quote-part de la promotion", SAISIE_F, EUROS),
    ]
    for i, (lib, val, com, fill, f) in enumerate(hyp, 5):
        _cellule(ws, f"A{i}", lib)
        _cellule(ws, f"B{i}", val, fill, f)
        _cellule(ws, f"C{i}", com)
    _entete(ws, 18, ["Résultat", "Montant", "Formule"])
    res = [
        ("Chiffre d'affaires", "=B5*B6", "Stagiaires × prix"),
        ("Location de salle", "=B8*B9", "Journées × prix par jour"),
        ("Formateurs externes", "=B10*B11", "Heures × coût horaire"),
        ("Matériel informatique", "=B12", ""),
        ("Examens", "=B5*B13", "Stagiaires × coût par candidat"),
        ("Jury blanc", "=B14", ""),
        ("Communication", "=B15", ""),
        ("Frais généraux", "=B16", ""),
        ("Total des charges", "=SUM(B20:B26)", ""),
        ("Résultat de la promotion", "=B19-B27", "Avant rémunération du dirigeant"),
        ("Seuil de rentabilité (stagiaires)", '=IF(B6>0,ROUNDUP((B27-B5*B13)/(B6-B13),0),"")', "Effectif minimal pour couvrir les charges"),
        ("Coût de la salle par heure-stagiaire", '=IF(B5>0,B20/(B7*B5),"")', "Indicateur de pilotage"),
    ]
    for i, (lib, f, com) in enumerate(res, 19):
        _cellule(ws, f"A{i}", lib, gras=lib in ("Total des charges", "Résultat de la promotion", "Seuil de rentabilité (stagiaires)"))
        _cellule(ws, f"B{i}", f, CALC_F, "0" if "Seuil" in lib else ('#,##0.00 "€"' if "heure" in lib else EUROS))
        _cellule(ws, f"C{i}", com)
    # Trésorerie mensuelle
    _entete(ws, 33, ["Mois", "Jours en centre", "Encaissements", "Location salle", "Autres charges", "Solde du mois", "Solde cumulé"])
    mois = {}
    for d, lieu, *_ in jours:
        cle = dt.date(d.year, d.month, 1)
        mois.setdefault(cle, 0)
        if lieu == "Centre":
            mois[cle] += 1
    cles = sorted(mois)
    ws["I33"], ws["I34"], ws["I35"], ws["I36"] = "Échéancier d'encaissement", "À l'entrée", "À mi-parcours", "À la fin"
    ws["J34"], ws["J35"], ws["J36"] = 0.3, 0.4, 0.3
    for ref in ("J34", "J35", "J36"):
        ws[ref].fill, ws[ref].number_format, ws[ref].font = SAISIE_F, "0%", NORMAL
    ws["I37"] = "Délai moyen de paiement (mois)"
    ws["J37"] = 2
    ws["J37"].fill, ws["J37"].font = SAISIE_F, NORMAL
    ws["I33"].font = GRAS
    milieu = len(cles) // 2
    for i, m in enumerate(cles):
        r = 34 + i
        _cellule(ws, f"A{r}", m, fmt="mmmm yyyy")
        _cellule(ws, f"B{r}", mois[m], fmt="0")
        # encaissement décalé du délai de paiement : simplifié par indices
        enc = []
        for k, part in ((0, "$J$34"), (milieu, "$J$35"), (len(cles) - 1, "$J$36")):
            enc.append(f'IF({i}={k}+$J$37,$B$19*{part},0)')
        _cellule(ws, f"C{r}", "=" + "+".join(enc), CALC_F, EUROS)
        _cellule(ws, f"D{r}", f"=-B{r}*$B$9", CALC_F, EUROS)
        _cellule(ws, f"E{r}", f"=-($B$21+$B$22+$B$23+$B$24+$B$25+$B$26)/{len(cles)}", CALC_F, EUROS)
        _cellule(ws, f"F{r}", f"=C{r}+D{r}+E{r}", CALC_F, EUROS)
        _cellule(ws, f"G{r}", f"=F{r}" if i == 0 else f"=G{r - 1}+F{r}", CALC_F, EUROS)
    for k in range(1, 3):
        r = 34 + len(cles) - 1 + k
        _cellule(ws, f"A{r}", f"Mois + {k} après la fin")
        _cellule(ws, f"B{r}", 0, fmt="0")
        i = len(cles) - 1 + k
        enc = [f'IF({i}={kk}+$J$37,$B$19*{part},0)' for kk, part in ((0, "$J$34"), (milieu, "$J$35"), (len(cles) - 1, "$J$36"))]
        _cellule(ws, f"C{r}", "=" + "+".join(enc), CALC_F, EUROS)
        _cellule(ws, f"D{r}", 0, CALC_F, EUROS)
        _cellule(ws, f"E{r}", 0, CALC_F, EUROS)
        _cellule(ws, f"F{r}", f"=C{r}+D{r}+E{r}", CALC_F, EUROS)
        _cellule(ws, f"G{r}", f"=G{r - 1}+F{r}", CALC_F, EUROS)
    for col, w in zip("ABCDEFGHIJ", (46, 16, 18, 18, 18, 18, 18, 3, 32, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"


def _feuille(ws, entetes, lignes, largeurs):
    _entete(ws, 1, entetes)
    for l in lignes:
        ws.append(l)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = NORMAL, BORD
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def classeur():
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan d'actions"
    _feuille(ws, ["Date", "Domaine", "Action", "Fait le", "Preuve (lien ou emplacement)"],
             [[x, dom, txt, "", ""] for x, dom, txt in jalons()], [14, 16, 90, 14, 30])
    for r in range(2, ws.max_row + 1):
        ws[f"A{r}"].number_format = "dd/mm/yyyy"
    for code in PROMOTIONS:
        _budget(wb, code)
    for code, p in PROMOTIONS.items():
        ws = wb.create_sheet(f"Heures {code}")
        jours, _ = planifier(code)
        lignes = []
        for d, lieu, m, k, titre, obj in jours:
            lignes.append([d, lieu, m["code"] if m else "PAE", titre, DUREE_JOUR, None, None])
        _feuille(ws, ["Date", "Lieu", "Module", "Séquence", "Heures prévues", "Heures réalisées", "Écart"], lignes, [14, 13, 9, 60, 14, 14, 10])
        for r in range(2, len(lignes) + 2):
            ws[f"A{r}"].number_format = "dd/mm/yyyy"
            ws[f"F{r}"].fill = SAISIE_F
            ws[f"G{r}"] = f'=IF(F{r}="","",F{r}-E{r})'
        tot = len(lignes) + 2
        ws[f"D{tot}"], ws[f"E{tot}"], ws[f"F{tot}"] = "TOTAL", f"=SUM(E2:E{tot - 1})", f"=SUM(F2:F{tot - 1})"
        for ref in (f"D{tot}", f"E{tot}", f"F{tot}"):
            ws[ref].font = GRAS
    ws = wb.create_sheet("Indicateurs")
    _feuille(ws, ["Indicateur", "MEM", "RPMS", "Formule ou source"], [
        ["Inscrits", None, None, "Saisie"], ["Entrés en formation", None, None, "Saisie"], ["Abandons", None, None, "Saisie"],
        ["Taux d'abandon", '=IF(B3="","",B4/B3)', '=IF(C3="","",C4/C3)', "Abandons / entrés"],
        ["Présentés à l'examen", None, None, "Saisie"], ["Titres complets obtenus", None, None, "Saisie"],
        ["Taux de réussite", '=IF(B6="","",B7/B6)', '=IF(C6="","",C7/C6)', "Titres / présentés"],
        ["Satisfaction moyenne (sur 4)", None, None, "Questionnaire EVAL-10"], ["En emploi à 6 mois", None, None, "Enquête"],
        ["Taux d'insertion à 6 mois", '=IF(B3="","",B10/(B3-B4))', '=IF(C3="","",C10/(C3-C4))', "En emploi / sortants"],
    ], [34, 14, 14, 30])
    for r in (5, 8, 11):
        for col in "BC":
            ws[f"{col}{r}"].number_format = "0%"
    ws = wb.create_sheet("Absences")
    _feuille(ws, ["Date", "Stagiaire (initiales)", "Promotion", "Type", "Justifiée ?", "Contact le", "Action", "Partenaire mobilisé"],
             [], [12, 18, 12, 14, 12, 12, 40, 20])
    out = SORTIE / DOSSIER / "PIL-03_Budget_tresorerie_tableau_de_bord.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def tout():
    return [doc_plan(), doc_pieces(), classeur()]


if __name__ == "__main__":
    for f in tout():
        print(f.name)
